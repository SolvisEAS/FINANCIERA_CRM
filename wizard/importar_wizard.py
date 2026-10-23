import base64
import io

import openpyxl

from odoo import fields, models
from odoo.exceptions import UserError

COLUMNAS = ['Nombre', 'CI', 'Teléfono', 'Email', 'Producto', 'Importe solicitado', 'Origen', 'Vendedora', 'Notas']
PRODUCTOS = {'prestamo': 'prestamo', 'préstamo': 'prestamo', 'tarjeta': 'tarjeta',
             'tarjeta de crédito': 'tarjeta', 'tarjeta de credito': 'tarjeta',
             'seguro': 'seguro', 'seguro vehicular': 'seguro'}
ORIGENES = {'meta': 'meta', 'meta lead ads': 'meta', 'whatsapp': 'whatsapp', 'importación': 'importacion',
            'importacion': 'importacion', 'manual': 'manual', 'referido': 'referido'}


def generar_plantilla():
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = 'Oportunidades'
    hoja.append(COLUMNAS)
    hoja.append(['Juan Pérez Benítez', '4123456', '0981 123 456', 'juan@ejemplo.com.py', 'Préstamo',
                 15000000, 'Importación', '', 'Pidió información por teléfono'])
    for celda in hoja[1]:
        celda.font = openpyxl.styles.Font(bold=True, color='FFFFFF')
        celda.fill = openpyxl.styles.PatternFill('solid', fgColor='0B5F7A')
    for columna, ancho in zip('ABCDEFGHI', [28, 12, 16, 26, 20, 18, 14, 16, 36]):
        hoja.column_dimensions[columna].width = ancho
    ayuda = libro.create_sheet('Ayuda')
    ayuda.append(['Producto: Préstamo, Tarjeta de crédito o Seguro vehicular'])
    ayuda.append(['Origen: Meta Lead Ads, WhatsApp, Importación, Manual o Referido (vacío = Importación)'])
    ayuda.append(['Vendedora: nombre o usuario. Vacío = se asigna por rotación'])
    ayuda.append(['Si el teléfono o email ya existe, se reutiliza el contacto y no se duplica la oportunidad'])
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


class FinancieraImportarWizard(models.TransientModel):
    _name = 'financiera.importar.wizard'
    _description = 'Importar oportunidades desde Excel'

    archivo = fields.Binary('Archivo (.xlsx)')
    nombre_archivo = fields.Char()
    resultado = fields.Text(readonly=True)

    def action_descargar_plantilla(self):
        return {'type': 'ir.actions.act_url', 'url': '/financiera/plantilla.xlsx', 'target': 'self'}

    def action_importar(self):
        self.ensure_one()
        if not self.archivo:
            raise UserError('Subí el archivo .xlsx con la plantilla.')
        try:
            libro = openpyxl.load_workbook(io.BytesIO(base64.b64decode(self.archivo)), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise UserError(f'No se pudo leer el archivo: {exc}') from exc
        hoja = libro.worksheets[0]
        filas = list(hoja.iter_rows(values_only=True))
        if not filas:
            raise UserError('El archivo está vacío.')
        encabezado = [str(c or '').strip().lower() for c in filas[0]]
        indice = {nombre.lower(): encabezado.index(nombre.lower()) for nombre in COLUMNAS if nombre.lower() in encabezado}
        if 'nombre' not in indice or 'teléfono' not in indice:
            raise UserError('Faltan columnas obligatorias: Nombre y Teléfono. Usá la plantilla.')
        Users = self.env['res.users'].sudo()
        Lead = self.env['crm.lead']
        creadas = existentes = errores = 0
        detalle = []
        for numero, fila in enumerate(filas[1:], start=2):
            def valor(col):
                pos = indice.get(col.lower())
                return fila[pos] if pos is not None and pos < len(fila) else None
            if not any(fila):
                continue
            nombre, telefono = valor('Nombre'), valor('Teléfono')
            if not nombre or not telefono:
                errores += 1
                detalle.append(f'Fila {numero}: falta nombre o teléfono')
                continue
            vendedora = Users.browse()
            if valor('Vendedora'):
                texto = str(valor('Vendedora')).strip()
                vendedora = Users.search(['|', ('login', '=ilike', texto), ('name', '=ilike', texto)], limit=1)
            try:
                with self.env.cr.savepoint():
                    lead, nueva = Lead._fin_ingresar({
                        'nombre': str(nombre).strip(), 'telefono': str(telefono), 'email': valor('Email'),
                        'ci': str(valor('CI') or '').strip() or False,
                        'producto': PRODUCTOS.get(str(valor('Producto') or '').strip().lower(), 'sin_definir'),
                        'importe': valor('Importe solicitado'),
                        'origen': ORIGENES.get(str(valor('Origen') or '').strip().lower(), 'importacion'),
                        'nota': valor('Notas'),
                        'user_id': vendedora.id or False,
                        'asignado_por': vendedora and 'Importación' or False,
                    })
            except Exception as exc:  # noqa: BLE001
                errores += 1
                detalle.append(f'Fila {numero}: {exc}')
                continue
            if nueva:
                creadas += 1
            else:
                existentes += 1
                detalle.append(f'Fila {numero}: {nombre} ya tenía una oportunidad abierta (se agregó una nota)')
        self.resultado = '\n'.join([f'Creadas: {creadas}', f'Ya existentes: {existentes}', f'Con error: {errores}', ''] + detalle)
        return {
            'type': 'ir.actions.act_window', 'res_model': self._name, 'res_id': self.id,
            'view_mode': 'form', 'target': 'new',
        }
